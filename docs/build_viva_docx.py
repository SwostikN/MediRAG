"""One-shot script to build docs/VIVA qsns.docx from a hard-coded question list."""
from docx import Document
from docx.shared import Pt, RGBColor

QUESTIONS = [
    # ── Section 1 — Basics & Motivation (1-15) ───────────────────────────
    ("Basics & Motivation", [
        "In one sentence, what is DocuMed AI and who is it built for?",
        "Why did you call it a \"health navigator\" rather than a medical diagnostic system?",
        "Why is Nepal the grounding context for this project rather than just a target market?",
        "What problem in Nepal's primary-care landscape does DocuMed AI try to solve?",
        "Who is your target user — patients, doctors, or both? Why?",
        "What is Retrieval-Augmented Generation (RAG), and why is it a better fit for this domain than fine-tuning?",
        "What is the difference between a chatbot that uses RAG and one that uses only an LLM?",
        "What does your system explicitly NOT do, and why are those refusals load-bearing?",
        "Why can't a general-purpose LLM like ChatGPT safely replace DocuMed AI for this use case?",
        "What is the corpus scope of DocuMed AI (which areas of medicine are in vs out of scope)?",
        "Why did you exclude surgery, oncology depth, ICU care, and specialist content from the corpus?",
        "What does the slogan \"clinical safety beats metrics\" mean in your context?",
        "How does your project respect Nepal's tiered health system (Health Post → PHCC → District → Tertiary)?",
        "What languages does the frontend support and why bilingual EN/NE?",
        "Walk us through what happens end-to-end when a user types: \"I have chest pain since yesterday morning\".",
    ]),
    # ── Section 2 — Architecture & Stack (16-30) ────────────────────────
    ("Architecture & Stack", [
        "Draw or describe the overall system architecture. What are the major components?",
        "Why did you choose FastAPI for the backend?",
        "Why Supabase with pgvector instead of a dedicated vector database like Pinecone or Weaviate?",
        "What is pgvector and how does it integrate with Postgres?",
        "What is MedCPT and why did you choose it over generic embedding models like OpenAI's text-embedding-3?",
        "Explain the difference between MedCPT's Article Encoder and Query Encoder. Why two models?",
        "What dimensionality are your embeddings, and how is the embedding stored in Postgres?",
        "Why is the Cohere Rerank model needed when you already have dense retrieval?",
        "What is the cross-encoder vs bi-encoder distinction, and where does each appear in your pipeline?",
        "Why do you use Groq for generation but Cohere for rerank — what's the rationale for the split?",
        "What is the fallback behaviour if Groq is unavailable?",
        "What is the cold-start startup hook in `RAG.py` doing, and why?",
        "Why is the NLI verification model loaded lazily rather than at import?",
        "What is the React frontend built with (Vite/TypeScript), and how is it served by FastAPI?",
        "Explain the role of `.env` variables — name three critical ones and what they control.",
    ]),
    # ── Section 3 — Retrieval Pipeline (31-45) ───────────────────────────
    ("Retrieval Pipeline", [
        "Walk through `_retrieve_ranked` step by step. What happens after the query is encoded?",
        "What is Reciprocal Rank Fusion (RRF) and why is `k=60` used? Cite the paper.",
        "How does your hybrid retrieval combine dense (cosine similarity) and lexical (BM25-ish tsvector) signals?",
        "What is the role of the `tsvector` column in the `chunks` table?",
        "What does `match_chunks_hybrid_filtered` filter on (country, authority, age, domain)?",
        "Walk through `build_filter()` — when does it apply stage- or domain-specific overrides?",
        "Explain the stage-weighted final score: `w_rerank * rerank + w_authority * authority + w_freshness * freshness`. Why do `intake` and `results` zero out authority + freshness?",
        "What is `RERANK_REFUSAL_THRESHOLD = 0.4` doing and why fail closed when it's not met?",
        "Why is the chunk pruning threshold `RERANK_CONTEXT_MIN = 0.2` distinct from the refusal threshold?",
        "What happens when Cohere Rerank fails — what is your synthetic fallback scoring scheme and why does it preserve gate semantics?",
        "Why do you re-rank session-uploaded chunks alongside corpus chunks before the LLM sees them?",
        "Why did you add a query rewrite step (`expand_for_retrieval`)? Give an example of the lay → clinical transformation.",
        "How do you handle conversational context for retrieval — what does `_retrieval_query_with_history` do?",
        "What is the `freshness_score`'s decay rate, and how does the authority tier (1-5) map to a score?",
        "Why do you de-duplicate sources for display but keep all chunks for the LLM context?",
    ]),
    # ── Section 4 — Stages & Routing (46-58) ─────────────────────────────
    ("Stages & Routing", [
        "What are the five stages in your state machine (intake, navigation, visit_prep, results, condition)?",
        "Explain the structured intake flow (Stage 1). Why 5 slot questions and what frameworks are they based on (SOCRATES / OPQRST / IMAI / PHQ-2)?",
        "Why is `intake_summary` framed as something the user reads TO a doctor, not as an assessment OF them?",
        "What is the pre-intake intent gate (`intent_gate.py`) and what four layers does it use?",
        "Walk through Layer-2 of the intent gate — why is the symptom-report detector intentionally greedy?",
        "What does Stage 2 (Navigation) produce, and how does it know which Nepal tier to recommend?",
        "Where do the urgent symptom and emergency overrides in `navigation.py` come from? Give two examples.",
        "Walk through Stage 4 (Results): how do you parse markers from a lab-report PDF?",
        "Why does `extract_lab_markers` use a two-pass strategy (same-line pass + windowed fallback)?",
        "Why do you only ever quote the reference range printed on the user's report rather than a \"standard\" range?",
        "What does the Stage 4 deterministic report-summary section add over the per-marker LLM blocks?",
        "What does the meta-question gate (`meta_question.py`) do, and why must it run BEFORE retrieval?",
        "What is the clarification stage and why does it NLI-verify the new answer against the prior assistant turn?",
    ]),
    # ── Section 5 — Safety Layers & Guardrails (59-75) ──────────────────
    ("Safety & Guardrails", [
        "What is the red-flag engine and where does it run in the pipeline?",
        "Why is the red-flag engine deterministic (YAML rules) rather than LLM-based?",
        "How does `redflag.py` handle nested `all_of` / `any_of` conditions?",
        "What is the difference between the red-flag engine, the refusal filter, and the scope guard?",
        "Explain the three-layer hallucination guardrail: claim classifier → NLI verifier → tier action.",
        "Why is the claim classifier (`guardrails.classify_claim`) deliberately regex/keyword-based instead of ML?",
        "What are the four \"clinical claim\" signals the classifier looks for (dose, threshold, diagnosis verb, duration)?",
        "Why are disclaimers skipped from NLI verification?",
        "Walk through the NLI verifier — what model do you use and how do you look up the entailment label index?",
        "Explain the three tier cutoffs: redact below 0.2, soften 0.2–0.5, keep above 0.5. Why is the hard-claim floor 0.5?",
        "What is the fusion-drift check (Phase-3B) catching that single-sentence NLI cannot?",
        "What is the inline-citation binding (Phase-3A), and why does it close the \"right claim, wrong citation\" failure?",
        "What does the scope-guard classifier do at the policy level that the regex refusal filter does not?",
        "Walk through `rewrite_emergency_numbers` — why does NHS \"call 999\" get rewritten to \"call 102\"?",
        "Why do you stream responses sentence-by-sentence (Option C) rather than streaming raw tokens?",
        "What is the emergency override in `classify_scope`, and why does it take precedence over diagnostic cues?",
        "When a hard-claim sentence (dose / diagnosis) fails NLI, why is it always redacted and never softened?",
    ]),
    # ── Section 6 — Database, RLS & Multi-Tenancy (76-83) ───────────────
    ("Database & Privacy", [
        "Walk through the `documents`, `chunks`, `citations`, `session_documents`, `session_chunks` tables. What is shared corpus vs session-private?",
        "Why do `/upload_pdf` (admin) and `/upload` (user) write to different tables? What invariant does that preserve?",
        "What Row Level Security (RLS) policies protect chat messages, sessions, and lab markers?",
        "How is `user_lab_markers` linked to a user and what happens on delete?",
        "Why is the `chunks_embedding_idx` an ivfflat index with `lists=100`? What is the trade-off vs HNSW?",
        "What index supports the lexical side of hybrid retrieval (`chunks_tsv_idx`)?",
        "How does the `query_log` table support auditability of every answer (prompt hash, NLI scores, refusal reasons)?",
        "What does the `documents.retracted` flag do, and why is it filtered in the RLS read policy and the SQL function?",
    ]),
    # ── Section 7 — Ingestion, Frontend & Uploads (84-90) ───────────────
    ("Ingestion & Frontend", [
        "Walk through the ingestion pipeline — `ingest/fetch.py`, `ingest/parse.py`, `ingest/run.py`, `ingest/sources/`.",
        "How does the document classifier (`document_classifier.py`) decide lab_report vs research_paper vs other?",
        "Why is the document classifier heuristic (regex) instead of an LLM call?",
        "What is the deduplication mechanism for re-uploaded PDFs (SHA-256 content hash)?",
        "What is the admin upload token (`ADMIN_UPLOAD_TOKEN`) protecting, and why does the endpoint fail closed when unset?",
        "Explain the `/upload/resolve` endpoint — why does the user need a disambiguation button for \"other\" uploads?",
        "How does the frontend render lab markers (table data) separately from the LLM-generated explainer prose, and why is that important on a refusal?",
    ]),
    # ── Section 8 — Evaluation & Empirical Work (91-95) ─────────────────
    ("Evaluation", [
        "Walk through your evaluation harness. What are the gold sets and what does each measure (must_refuse, coverage, hallucination, recall@k, stage2, stage4)?",
        "What metric do you use for hallucination scoring, and how do you compare two configurations (`score_hallucination_compare.py`)?",
        "What is `score_must_refuse.py` checking, and why is over-refusal a failure mode too?",
        "How do you A/B test inline citations and fusion-drift via the `INLINE_CITATIONS_ENABLED` and `FUSION_DRIFT_ENABLED` flags?",
        "What was the Stage-2 Cohere-fallback issue on 2026-04-20, and how did the Groq re-run change the baseline?",
    ]),
    # ── Section 9 — Critical / Research-Style (96-100) ──────────────────
    ("Critical & Forward-Looking", [
        "What are the current limitations of your system — both technical and clinical?",
        "How would you scale DocuMed AI to handle 100k Nepalese users concurrently? Where would the bottleneck be?",
        "If you had to pick the single biggest risk of this system reaching production, what would it be and how would you mitigate it?",
        "Why did you choose a defensive, multi-layered guardrail design instead of trusting one strong LLM with a careful system prompt?",
        "If you had three more months, what would you build next and why — corpus expansion, Retraction Watch sync, longitudinal lab tracking, or something else?",
    ]),
]


def main() -> None:
    doc = Document()

    # Title
    title = doc.add_heading("DocuMed AI — Viva Questions", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6B)

    sub = doc.add_paragraph()
    sub_run = sub.add_run(
        "100 questions, simple → hard, covering the full system — "
        "architecture, retrieval, safety layers, database, evaluation, "
        "and design trade-offs."
    )
    sub_run.italic = True
    sub_run.font.size = Pt(11)

    doc.add_paragraph(
        "Tip: for every question, anchor your answer in (a) what the code does, "
        "(b) why you chose that design, and (c) the trade-off / failure mode it "
        "addresses. Supervisors reward reasoning over recall."
    )

    counter = 1
    for section_title, questions in QUESTIONS:
        h = doc.add_heading(section_title, level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6B)
        for q in questions:
            p = doc.add_paragraph()
            num = p.add_run(f"Q{counter}. ")
            num.bold = True
            p.add_run(q)
            counter += 1
        doc.add_paragraph()  # blank line between sections

    out_path = "/Users/rewati/Desktop/MediRAG/docs/VIVA qsns.docx"
    doc.save(out_path)
    print(f"Wrote {out_path} with {counter - 1} questions across {len(QUESTIONS)} sections.")


if __name__ == "__main__":
    main()
